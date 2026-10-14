"""MarketLens - Data Exporter Module.

Handles exporting cleaned market datasets into CSV and professional Excel (.xlsx) formats.
Applies automated styling, custom column widths, number formats, and freeze panes
using openpyxl and pandas for consulting-ready deliverables.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import logging
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

logger = logging.getLogger("marketlens.exporter")

# Corporate consulting color palette
HEADER_FILL_COLOR = "1F4E79"   # Navy Blue
HEADER_FONT_COLOR = "FFFFFF"   # White
ZEBRA_FILL_COLOR = "F9FAFB"    # Subtle light gray for alternating rows
BORDER_COLOR = "D9D9D9"        # Soft border gray


def ensure_dataframe(data: Union[List[Dict[str, Any]], pd.DataFrame]) -> pd.DataFrame:
    """Convert input records into a pandas DataFrame safely."""
    if isinstance(data, pd.DataFrame):
        return data.copy()
    if not data:
        return pd.DataFrame()
    return pd.DataFrame(data)


def export_to_csv(
    data: Union[List[Dict[str, Any]], pd.DataFrame],
    output_path: Union[str, Path],
    encoding: str = "utf-8-sig",
) -> Optional[Path]:
    """Export dataset to a CSV file with standardized encoding."""
    df = ensure_dataframe(data)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    if df.empty:
        logger.warning("Empty dataset provided for CSV export to: %s", target)
        df.to_csv(target, index=False, encoding=encoding)
        return target

    df.to_csv(target, index=False, encoding=encoding)
    logger.info("Exported CSV successfully: %s (%d rows)", target, len(df))
    return target


def format_excel_worksheet(worksheet: Any, df: pd.DataFrame) -> None:
    """Apply professional corporate styling to an openpyxl worksheet."""
    header_font = Font(name="Calibri", size=11, bold=True, color=HEADER_FONT_COLOR)
    header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    zebra_fill = PatternFill(start_color=ZEBRA_FILL_COLOR, end_color=ZEBRA_FILL_COLOR, fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    center_alignment = Alignment(horizontal="center", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")
    right_alignment = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )

    # 1. Format Header Row (Row 1)
    worksheet.row_dimensions[1].height = 26
    for col_idx in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 2. Format Data Rows & Set Column Types
    for row_idx in range(2, len(df) + 2):
        worksheet.row_dimensions[row_idx].height = 20
        is_even_row = (row_idx % 2 == 0)

        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border

            if not is_even_row:
                cell.fill = zebra_fill

            # Apply column-specific formatting
            if "price" in col_name.lower():
                cell.number_format = "£#,##0.00"
                cell.alignment = right_alignment
            elif "rating" in col_name.lower():
                cell.number_format = "0"
                cell.alignment = center_alignment
            elif col_name.lower() in ("is_in_stock", "in_stock"):
                cell.alignment = center_alignment
            elif "url" in col_name.lower() or "timestamp" in col_name.lower():
                cell.alignment = left_alignment
            else:
                cell.alignment = left_alignment

    # 3. Auto-adjust column widths based on maximum string length
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for val in df[col_name]:
            if pd.notna(val):
                val_len = len(str(val))
                if val_len > max_len:
                    max_len = val_len

        fitted_width = min(max(max_len + 4, 12), 55)
        worksheet.column_dimensions[col_letter].width = fitted_width

    # 4. Freeze Header Panes so Row 1 stays locked when scrolling
    worksheet.freeze_panes = "A2"

    # 5. Ensure Gridlines are displayed in Excel
    worksheet.views.sheetView[0].showGridLines = True


def export_to_excel(
    data: Union[List[Dict[str, Any]], pd.DataFrame],
    output_path: Union[str, Path],
    sheet_name: str = "Market Data",
) -> Optional[Path]:
    """Export dataset to an Excel workbook with styling and auto-widths."""
    df = ensure_dataframe(data)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    if df.empty:
        logger.warning("Empty dataset provided for Excel export to: %s", target)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        wb.save(target)
        return target

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.sheets[sheet_name]
        format_excel_worksheet(worksheet, df)

    logger.info("Exported Excel workbook successfully: %s (%d rows)", target, len(df))
    return target


def export_dataset(
    data: Union[List[Dict[str, Any]], pd.DataFrame],
    base_output_path: Union[str, Path],
    formats: Optional[List[str]] = None,
) -> Dict[str, Path]:
    """Export data to both CSV and Excel simultaneously."""
    if formats is None:
        formats = ["csv", "xlsx"]

    base = Path(base_output_path)
    if base.suffix in (".csv", ".xlsx"):
        base = base.with_suffix("")

    results: Dict[str, Path] = {}

    if "csv" in formats:
        csv_path = base.with_suffix(".csv")
        export_to_csv(data, csv_path)
        results["csv"] = csv_path

    if "xlsx" in formats:
        excel_path = base.with_suffix(".xlsx")
        export_to_excel(data, excel_path)
        results["xlsx"] = excel_path

    return results

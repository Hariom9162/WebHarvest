"""MarketLens - Exporter Unit Tests.

Tests CSV encoding, Excel formatting, freeze panes, auto widths,
and safe handling of empty datasets using pytest temporary directories.
"""

from pathlib import Path
import openpyxl
import pandas as pd
import pytest
from marketlens.exporter import export_dataset, export_to_csv, export_to_excel

SAMPLE_DATA = [
    {
        "title": "Clean Code",
        "price_gbp": 32.50,
        "rating_stars": 4,
        "is_in_stock": True,
        "availability": "In Stock",
        "product_url": "https://books.toscrape.com/clean-code.html",
        "source_url": "https://books.toscrape.com/",
        "extraction_timestamp": "2026-10-15T00:00:00Z",
    },
    {
        "title": "The Pragmatic Programmer",
        "price_gbp": 45.00,
        "rating_stars": 5,
        "is_in_stock": True,
        "availability": "In Stock",
        "product_url": "https://books.toscrape.com/pragmatic.html",
        "source_url": "https://books.toscrape.com/",
        "extraction_timestamp": "2026-10-15T00:00:00Z",
    },
]


def test_export_to_csv(tmp_path: Path) -> None:
    """Test CSV export with row count and column verification."""
    csv_file = tmp_path / "test_books.csv"
    saved_path = export_to_csv(SAMPLE_DATA, csv_file)
    assert saved_path.exists()

    df = pd.read_csv(saved_path, encoding="utf-8-sig")
    assert len(df) == 2
    assert "price_gbp" in df.columns
    assert df.loc[0, "title"] == "Clean Code"


def test_export_to_excel_styling(tmp_path: Path) -> None:
    """Test Excel export with openpyxl verification of freeze panes and column widths."""
    xlsx_file = tmp_path / "test_books.xlsx"
    saved_path = export_to_excel(SAMPLE_DATA, xlsx_file, sheet_name="TestSheet")
    assert saved_path.exists()

    wb = openpyxl.load_workbook(saved_path)
    assert "TestSheet" in wb.sheetnames
    ws = wb["TestSheet"]

    # Verify freeze pane is on row 2
    assert ws.freeze_panes == "A2"

    # Verify 2 data rows + 1 header = 3 rows
    assert ws.max_row == 3

    # Verify column A has adjusted width > 15
    assert ws.column_dimensions["A"].width > 15

    # Verify price formatting
    price_cell = ws["B2"]
    assert price_cell.number_format == "£#,##0.00"


def test_export_empty_dataset(tmp_path: Path) -> None:
    """Test that exporting an empty dataset doesn't crash and generates a valid empty file."""
    empty_csv = tmp_path / "empty.csv"
    empty_xlsx = tmp_path / "empty.xlsx"

    csv_res = export_to_csv([], empty_csv)
    xlsx_res = export_to_excel([], empty_xlsx)

    assert csv_res.exists()
    assert xlsx_res.exists()


def test_export_dataset_both_formats(tmp_path: Path) -> None:
    """Test simultaneous export of both CSV and XLSX."""
    base_path = tmp_path / "output_package"
    results = export_dataset(SAMPLE_DATA, base_path, formats=["csv", "xlsx"])

    assert "csv" in results
    assert "xlsx" in results
    assert results["csv"].exists()
    assert results["xlsx"].exists()
